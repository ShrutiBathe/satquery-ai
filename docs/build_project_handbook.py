from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = Path(__file__).resolve().parent / "SatQuery_AI_Project_Handbook.docx"

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "EAF3F8"
PALE = "F5F7FA"
GRID = "D9D9D9"
TEXT = "1F2937"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_text(paragraph, text, bold=False, color=TEXT, size=10.5, italic=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_paragraph(doc, text="", bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, bold=True)
        add_text(p, text[len(bold_lead):])
    else:
        add_text(p, text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        add_text(p, item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, head in enumerate(headers):
        cell = hdr.cells[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margin(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_text(p, head, bold=True, color="FFFFFF", size=9.2)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else PALE)
            set_cell_border(cell)
            set_cell_margin(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            add_text(p, str(value), size=8.8)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    add_text(p, text, bold=True, color="000000", size=15 if level == 1 else 12)
    return p


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F6F9")
    set_cell_border(cell, "B8C2CC")
    set_cell_margin(cell, 120, 140, 120, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run.font.size = Pt(8.7)
        run.font.color.rgb = RGBColor.from_string("263238")
        if index < len(lines) - 1:
            run.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Title"].font.size = Pt(27)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(footer, "SatQuery AI Project Handbook  |  ", color="5B6573", size=8)
    page_number(footer)


def build():
    doc = Document()
    configure(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(title, "SatQuery AI Project Handbook", bold=True, color="000000", size=27)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(subtitle, "Current Implementation Architecture Team Guide and Evaluator Brief", color="3F4B59", size=13)
    doc.add_paragraph()
    add_paragraph(doc, "This handbook explains the current SatQuery AI system as implemented in the repository. It is designed for all six team members and for evaluator conversations. It separates working capabilities from planned capabilities so that the team can demonstrate the system accurately and describe the roadmap confidently.")
    add_paragraph(doc, "Current conclusion: SatQuery AI already has a working Streamlit experience, a live FastAPI backend, LangGraph based request routing, input validation, GeoTIFF aware preprocessing, structured API responses, and a complete demo simulation experience. The specialist AI models and deployment infrastructure are interface-ready but are not yet integrated as real inference engines.")

    heading(doc, "How to use this handbook", 1)
    add_bullets(doc, [
        "Use Sections 1 to 5 to understand the product and end to end data flow.",
        "Use Sections 6 to 11 to understand each teammate module and its current boundary.",
        "Use Sections 12 to 14 before a demonstration or evaluator review.",
        "Treat the status labels carefully: implemented means present and exercised; planned means designed but not yet connected to real model inference.",
    ])

    heading(doc, "Contents", 1)
    contents = [
        "1. Product purpose and user value", "2. Current implementation snapshot", "3. Repository and team module map",
        "4. System architecture", "5. End to end data flow", "6. Frontend module", "7. Backend module",
        "8. AI agent module", "9. Preprocessing module", "10. AI models module", "11. Infrastructure module",
        "12. API contract and data artifacts", "13. Testing and demonstration", "14. Remaining work and roadmap",
        "15. Evaluator explanation script", "16. Team operating checklist",
    ]
    add_bullets(doc, contents)

    doc.add_page_break()
    heading(doc, "1 Product purpose and user value", 1)
    add_paragraph(doc, "SatQuery AI is an interactive assistant for multimodal remote sensing analysis. A user supplies one satellite scene or a pair of scenes, writes a question in ordinary language, and receives a routed analysis result with visual evidence, a natural language answer, metrics, confidence information, and an execution trace.")
    add_paragraph(doc, "The product is aimed at researchers, disaster response teams, urban planners, and users who are not expected to construct GIS pipelines manually. The intended tasks are visual question answering, object or region grounding, bi temporal change detection, and optical plus SAR analysis.")
    add_table(doc, ["User need", "SatQuery response", "Current status"], [
        ["Ask what is in one scene", "Routes to VQA and returns a standardized answer payload", "Routing and response flow implemented; real VQA model pending"],
        ["Locate assets such as buildings or tanks", "Routes to grounding and supports evidence outputs", "Routing and response flow implemented; real grounding model pending"],
        ["Compare before and after imagery", "Routes to change detection and aligns GeoTIFF inputs", "Preprocessing implemented; real change model pending"],
        ["Use optical and SAR jointly", "Routes to optical SAR and prepares paired inputs", "Pair preprocessing implemented; real fusion model pending"],
        ["Evaluate the product quickly", "Runs four procedural satellite scenarios in Demo Mode", "Implemented"],
    ], [1.55, 3.45, 1.75])

    heading(doc, "2 Current implementation snapshot", 1)
    add_table(doc, ["Capability", "Status", "Evidence in repository"], [
        ["Streamlit application", "Implemented", "frontend app, dashboard, workspace, uploads, result views, history, exports"],
        ["Demo simulation", "Implemented", "mock_service and sample_data create scenarios, overlays, mock metrics, and follow ups"],
        ["FastAPI API", "Implemented", "health endpoint, versioned analyze endpoint, upload storage, CORS, static output mount"],
        ["Backend orchestration", "Implemented", "backend service validates, routes, preprocesses, dispatches, and assembles results"],
        ["LangGraph routing", "Implemented and connected", "agent graph is invoked through agent.router from backend.service"],
        ["Geo validation and preprocessing", "Implemented", "format checks, RGB selection, resize, GeoTIFF alignment, paired preparation"],
        ["Specialist model inference", "Pending", "backend adapters intentionally return readiness placeholder results"],
        ["Production deployment stack", "Pending", "infrastructure directory and docker compose currently contain no implemented deployment definition"],
    ], [2.0, 1.25, 3.5])
    add_paragraph(doc, "Important distinction for demonstrations: Demo Mode produces rich synthetic evidence and confidence values. Live Mode calls the backend and confirms the real pipeline plumbing. At present, Live Mode returns a clear model integration pending answer after successful routing and preprocessing. The team must not claim that the demo overlays are generated by deployed neural models.", bold_lead="Important distinction for demonstrations: ")

    doc.add_page_break()
    heading(doc, "3 Repository and team module map", 1)
    add_table(doc, ["Team module", "Repository area", "Primary responsibility", "Current integration point"], [
        ["Frontend", "frontend", "Streamlit user experience, upload workflow, presentation, demo mode", "Calls services.api_client"],
        ["Backend", "backend", "API contract, validation, orchestration, error handling, response assembly", "Calls agent, geo, and future models"],
        ["AI agent", "agent", "Intent classification and task routing using LangGraph", "Called by backend.service"],
        ["Preprocessing", "geo", "Image validation, TIFF handling, band selection, resize, geospatial alignment", "Called by backend.service"],
        ["AI models", "models", "Real VQA, grounding, change, and fusion inference", "To be called by backend specialist adapters"],
        ["Infrastructure", "infrastructure and docker compose", "Packaging, reproducible environments, deployment, observability", "To host frontend and backend"],
    ], [1.1, 1.25, 2.55, 1.85])
    add_paragraph(doc, "The backend is the integration boundary. It should not reimplement the agent classifier, raster preprocessing, or model logic. It accepts a stable request, asks the agent for the task, asks preprocessing for model ready paths, calls the correct model adapter, and returns one stable AnalysisResult object to the frontend.")

    heading(doc, "4 System architecture", 1)
    add_code(doc, [
        "User", "  -> Streamlit Frontend", "       -> Demo Mode: mock service and procedural evidence", "       -> Live Mode: POST /api/v1/analyze", "            -> FastAPI upload and metadata handling", "            -> Backend analysis service", "                 -> Agent router and LangGraph task decision", "                 -> Geo validation and preprocessing", "                 -> Specialist adapter", "                      -> Future real AI model", "                 -> AnalysisResult JSON", "            -> Frontend result viewer and execution trace",
    ])
    add_paragraph(doc, "There are two intentionally different paths. Demo Mode is a self contained presentation and UX validation path. Live Mode is the integration path and is the correct path for validating that frontend, backend, agent, preprocessing, and future models work together.")

    heading(doc, "5 End to end data flow", 1)
    heading(doc, "5 1 Live mode request lifecycle", 2)
    add_table(doc, ["Step", "Component", "Action", "Output"], [
        ["1", "Frontend upload", "User selects Image A and optionally Image B, enters a query, and selects a mode", "PIL images, names, and metadata held in Streamlit session state"],
        ["2", "API client", "Converts PIL images to PNG bytes and sends multipart form data", "POST body with query, mode, metadata, image_a, image_b"],
        ["3", "FastAPI", "Checks query, allowed mode, filename extension, saves uploads into a unique runtime directory", "AnalysisRequest with server side image paths"],
        ["4", "Backend service", "Validates request shape and invokes the agent router", "Task, required image count, routing reason"],
        ["5", "LangGraph agent", "Classifies the natural language query and validates required images", "vqa, grounding, change_detection, or optical_sar"],
        ["6", "Geo module", "Validates files and prepares model ready inputs", "Processed paths plus preprocessing metadata"],
        ["7", "Specialist adapter", "Receives processed inputs and query", "Current placeholder result; future model output"],
        ["8", "Result assembler", "Normalizes trace, evidence, metrics, confidence, and metadata", "AnalysisResult JSON"],
        ["9", "Frontend", "Normalizes response, stores it in session state, displays results and history", "Visible analysis workspace"],
    ], [0.42, 1.2, 3.3, 2.0])
    heading(doc, "5 2 Demo mode lifecycle", 2)
    add_paragraph(doc, "Demo Mode never calls FastAPI. services.mock_service selects a workflow using frontend heuristics, and services.sample_data generates one of four procedural scenarios: urban change, industrial tank grounding, flood inundation with optical SAR, or maritime VQA. It creates presentation quality evidence layers such as heatmaps, annotated boxes, saliency overlays, answers, confidence scores, and traces. This is valuable for product demonstration, but it is simulation rather than live model inference.")

    doc.add_page_break()
    heading(doc, "6 Frontend module", 1)
    add_paragraph(doc, "The frontend is a Python Streamlit application with a dashboard and an analysis workspace. Its entry point is frontend/app.py. It keeps interaction state in Streamlit session state so navigation and results survive Streamlit reruns.")
    add_table(doc, ["Area", "What it does", "Integration boundary"], [
        ["app.py", "Configures Streamlit, injects theme CSS, renders the sidebar and routes pages", "Imports components and session state helpers"],
        ["components upload and query input", "Collects imagery, dates, metadata, selected mode, and natural language query", "Supplies inputs to workflow"],
        ["components workflow", "Calls client.analyze, displays staged pipeline animation, saves result in session history", "Consumes normalized backend result"],
        ["services api_client", "Chooses mock service in Demo Mode or HTTP request in Live Mode", "Calls FastAPI /api/v1/analyze"],
        ["result components", "Render visual evidence, answer, confidence, trace, follow ups, exports", "Read session state only"],
    ], [1.55, 3.5, 1.9])
    add_paragraph(doc, "Live mode caveat: api_client serializes every uploaded PIL image as PNG. That permits standard single image flows but loses the original GeoTIFF format and georeferencing. Since current backend change detection requires two GeoTIFF files, an end to end live change detection request will fail until the upload transport preserves original TIFF uploads or the preprocessing contract changes.", bold_lead="Live mode caveat: ")

    heading(doc, "7 Backend module", 1)
    add_paragraph(doc, "The backend is the application integration layer. backend/main.py hosts FastAPI. It exposes GET /health, POST /api/v1/analyze, and a backward compatible POST /analyze alias. The analysis endpoint receives multipart form data, validates inputs, writes uploads below backend/runtime/uploads/ANL-..., builds an AnalysisRequest, and calls backend.service.analyze.")
    add_table(doc, ["Backend file", "Responsibility"], [
        ["main.py", "FastAPI app, CORS, static output mount, health, upload validation, multipart parsing, file saving"],
        ["schemas.py", "Pydantic contracts: AnalysisRequest, AnalysisResult, DetectedTask, EvidenceItem, Metric, TraceItem"],
        ["service.py", "Main orchestration: agent route, geo validation, preprocessing, specialist dispatch, trace and result assembly"],
        ["validator.py", "Reusable request validation helper. It exists but is not currently invoked by service.analyze"],
        ["errors.py", "SatQueryError base plus ValidationError and ModelError"],
    ], [1.45, 5.5])
    add_paragraph(doc, "The backend to agent connection is explicit. backend/service.py imports agent.router.understand_query as run_agent_router. Its understand_query adapter passes query, image_count, and metadata to the agent. ValueError from agent validation is converted to backend ValidationError so client correctable errors are reported correctly by the API.")

    heading(doc, "Backend specialist adapter behavior today", 2)
    add_table(doc, ["Task", "Current adapter", "Today", "Required next integration"], [
        ["vqa", "run_vqa", "Returns a successful readiness payload after preprocessing", "Call the VQA model and return model answer, confidence, evidence"],
        ["grounding", "run_grounding", "Returns a successful readiness payload after preprocessing", "Call detector or grounding model and return boxes or masks"],
        ["change_detection", "run_change_detection", "Returns readiness payload after paired GeoTIFF alignment", "Run change model and emit change mask, heatmap, statistics"],
        ["optical_sar", "run_optical_sar", "Returns readiness payload after paired preparation", "Run fusion model and emit fused evidence and metrics"],
    ], [1.2, 1.45, 2.2, 2.1])

    doc.add_page_break()
    heading(doc, "8 AI agent module", 1)
    add_paragraph(doc, "The agent module is a lightweight deterministic routing agent implemented with LangGraph. agent/query_understanding.py examines the lowercased query and chooses a task. agent/graph.py wraps that classifier in a StateGraph. agent/router.py is the public backend entry point and validates that enough images were provided for the selected task.")
    add_table(doc, ["Task", "Agent triggers", "Required images", "Reason returned"], [
        ["optical_sar", "Query includes optical and SAR or optical SAR", "2", "Joint optical and SAR analysis requested"],
        ["change_detection", "change, changed, difference, compare, before and after, between two images", "2", "Temporal or comparative analysis requested"],
        ["grounding", "where, find, locate, highlight, mark, region, bounding", "1", "User wants spatial location or highlighting"],
        ["vqa", "Fallback for all other questions", "1", "General image analysis request"],
    ], [1.4, 3.15, 0.85, 1.55])
    add_paragraph(doc, "The LangGraph state contains query, image_count, metadata, task, required_images, and reason. The graph begins at understand_query, calls the classifier, validates the output fields, then uses conditional edges to end in the selected task. At this stage, the graph performs routing only; specialist model nodes are intentionally outside the graph and controlled by the backend.")
    add_paragraph(doc, "Current agent limitation: routing is keyword based, not LLM or embedding based. It is suitable for a reliable prototype routing demonstration. A future version can replace only agent/query_understanding.py or add an LLM classifier while retaining the backend contract.")

    heading(doc, "9 Preprocessing module", 1)
    add_paragraph(doc, "The geo module owns image preparation. It accepts JPG, JPEG, PNG, TIF, and TIFF. It validates file existence and task specific image counts, reads ordinary images through Pillow and GeoTIFF imagery through rasterio, then returns processed output paths and metadata.")
    add_table(doc, ["Task path", "Preprocessing behavior", "Output"], [
        ["VQA or grounding", "Reads image; resizes to target 100 by 150; for multiband TIFF selects bands 1, 2, 3; writes PNG", "outputs/processed/image1.png"],
        ["Change detection", "Requires two TIFFs; reprojects the after image to the before image grid using bilinear resampling", "outputs/processed/before.tif and after.tif"],
        ["Optical SAR", "Reads both inputs; resizes each to common target; saves TIFF or PNG according to input type", "outputs/processed/optical and sar"],
    ], [1.4, 3.7, 1.85])
    add_paragraph(doc, "Geospatial alignment is meaningful: for change detection, the post event raster is resampled onto the baseline raster grid using the baseline CRS, transform, width, and height. This avoids comparing pixels that represent different geographic locations. Model owners should keep this guarantee in mind when they consume the two processed change images.")

    heading(doc, "10 AI models module", 1)
    add_paragraph(doc, "The models directory currently has no model implementation files. The model team must provide callable inference functions or service clients that the backend adapters can invoke. The existing adapter signatures define the handoff shape: a processed image path and query for VQA or grounding; two processed image paths for change detection; optical path, SAR path, and query for fusion.")
    add_table(doc, ["Model capability", "Expected input", "Expected output for backend"], [
        ["VQA", "RGB model ready image path and query", "answer text, confidence, optional attention map, metrics, trace"],
        ["Grounding", "RGB model ready image path and query", "boxes, masks, labels, per detection confidence, annotated overlay"],
        ["Change detection", "Aligned before and after raster paths", "binary or continuous change mask, heatmap, changed area statistics"],
        ["Optical SAR fusion", "Paired optical and SAR paths and query", "fused prediction, water or flood mask, sensor contribution metrics"],
    ], [1.5, 2.35, 3.1])
    add_paragraph(doc, "Model contract recommendation: each model adapter should return a dictionary with success, answer, confidence, confidence_breakdown, summary_bullets, metrics, evidence, visual_evidence, output_paths, statistics, and trace. This matches the current backend result assembly and prevents frontend redesign.")

    doc.add_page_break()
    heading(doc, "11 Infrastructure module", 1)
    add_paragraph(doc, "Infrastructure is not yet implemented in the repository. docker-compose.yml is currently empty, and the infrastructure directory contains no deployment files. The application can be run locally with FastAPI and Streamlit, but it does not yet have a reproducible container build, service orchestration, persistent artifact storage, secret management, logging stack, or deployment target.")
    add_table(doc, ["Need", "Recommended deliverable", "Why it matters"], [
        ["Dependency reproducibility", "requirements files or pyproject with locked versions", "Every teammate and evaluator can run the same environment"],
        ["Containers", "Dockerfile for backend and frontend plus docker compose services", "One command local startup and consistent deployment"],
        ["Artifact storage", "Mounted volume or object store for uploads and evidence", "Avoid losing results and manage disk growth"],
        ["Configuration", "Environment variables for backend URL, model locations, storage", "Separate local, demo, and deployment settings"],
        ["Observability", "Structured logs, request IDs, health and readiness checks", "Diagnose model and pipeline failures"],
        ["Security", "Restrictive CORS, upload size limits, file scanning policy, auth when exposed", "Safe public deployment"],
    ], [1.5, 3.35, 2.1])

    heading(doc, "12 API contract and data artifacts", 1)
    heading(doc, "12 1 Analyze request", 2)
    add_code(doc, [
        "POST /api/v1/analyze", "multipart/form-data", "query: string", "mode: auto_detect | single_image | compare_images | optical_sar", "metadata_a: JSON object string", "metadata_b: JSON object string", "image_a: file required", "image_b: file optional",
    ])
    heading(doc, "12 2 Result structure", 2)
    add_code(doc, [
        "AnalysisResult", "  success, status, analysis_id, query, mode, task", "  detected_task: task_id, title, reason, required_images", "  answer, confidence, confidence_breakdown, summary_bullets", "  metrics, evidence, visual_evidence, output_paths, statistics", "  trace, analysis_trace, error",
    ])
    add_paragraph(doc, "The API serves generated evidence below /outputs. Uploaded files are stored below backend/runtime/uploads, while the configured evidence directory is backend/runtime/evidence. Processed files produced by the current geo module are written below outputs/processed relative to the working directory. Infrastructure work should consolidate these locations and add cleanup policies.")

    heading(doc, "13 Testing and demonstration", 1)
    add_table(doc, ["Test", "What was verified", "Result"], [
        ["Agent routing unit test", "VQA, grounding, change detection, optical SAR, and one image validation failure", "Passed using unittest"],
        ["Live frontend grounding test", "Demo Mode off; Streamlit sent a prepared single scene to FastAPI; backend called agent and preprocessing; frontend displayed result", "Passed"],
        ["Backend health test", "GET /health returned service status and version", "Passed"],
        ["Geo tests", "Existing tests reference local TIFF fixtures that are absent from repository", "Need fixture setup before suite can pass"],
    ], [1.65, 3.7, 1.55])
    heading(doc, "Recommended evaluator demo", 2)
    add_bullets(doc, [
        "Start the FastAPI backend and Streamlit frontend using the same Python environment.",
        "Open the Streamlit application and first show Demo Mode for polished product scenarios.",
        "Explain that Demo Mode validates the user experience with procedural satellite scenes.",
        "Turn Demo Mode off, load the single image grounding benchmark, and use a query such as Where are the fuel storage tanks and warehouse facilities.",
        "Show the AI Backend Connected indicator, the completed result, and the routing reason in the execution trace.",
        "State that the live model adapter is presently a scaffold. The important verified live result is that the request passed through frontend, FastAPI, LangGraph routing, geo preprocessing, and standardized response assembly.",
    ])

    doc.add_page_break()
    heading(doc, "14 Remaining work and roadmap", 1)
    add_table(doc, ["Priority", "Owner module", "Work item", "Completion criterion"], [
        ["P0", "AI models", "Implement and expose one real specialist model, ideally grounding or VQA first", "Live response contains real model evidence and confidence"],
        ["P0", "Frontend and backend", "Preserve original GeoTIFF bytes and metadata in live upload flow", "Live change detection accepts two TIFFs end to end"],
        ["P0", "Infrastructure", "Create reproducible requirements, Dockerfiles, and docker compose", "New machine starts the system with documented command"],
        ["P1", "Backend", "Integrate model functions into current specialist adapters", "No placeholder answer for integrated task"],
        ["P1", "Geo and models", "Agree normalization, band order, nodata handling, and output formats", "Model input contract documented and regression tested"],
        ["P1", "Agent", "Expand routing beyond keywords using metadata and potentially an LLM classifier", "Ambiguous multi intent queries resolve predictably"],
        ["P1", "All modules", "Add real raster fixtures and end to end API tests", "CI test suite runs without external manual files"],
        ["P2", "Backend and infrastructure", "Add artifact lifecycle, request IDs, structured logs, and secure CORS", "Deployment is observable and safer"],
    ], [0.55, 1.35, 3.25, 1.8])
    heading(doc, "Known technical gaps to explain honestly", 2)
    add_bullets(doc, [
        "The README describes aspirational neural capabilities more broadly than the current live backend. The handbook status table is the source of truth for implementation discussion.",
        "Frontend Demo Mode evidence is simulated and should not be described as results of deployed trained models.",
        "Live single image routing and preprocessing have been exercised; live model inference has not been implemented.",
        "The current backend mode field is accepted at the API layer, while task selection is owned by the agent query classifier. The team should agree whether explicit mode must override agent intent in a later change.",
        "agent/_init_.py is named differently from conventional __init__.py. Python namespace imports work in the present runtime, but renaming it is recommended before packaging.",
    ])

    heading(doc, "15 Evaluator explanation script", 1)
    add_paragraph(doc, "Start with the problem: Remote sensing analysis is powerful but often requires specialist GIS workflows. SatQuery AI makes the first interaction natural language. A user can upload a scene or a pair of scenes and ask a question such as where are the tanks, what changed, or map flood extent using optical and SAR data.")
    add_paragraph(doc, "Then explain the architecture: The Streamlit frontend collects imagery and a query. In Live Mode it sends a multipart request to FastAPI. The backend stores and validates files, calls our LangGraph agent for task selection, calls the geo preprocessing module for image preparation, dispatches to the matching specialist adapter, and returns a standardized response with trace and evidence fields.")
    add_paragraph(doc, "Then explain the current achievement: We have verified a real live frontend to backend to agent to preprocessing round trip for a grounding request. The application showed AI Backend Connected and returned a Visual Grounding routing result. The specialist model adapters are deliberately separated so the model team can plug in inference without changing the frontend contract.")
    add_paragraph(doc, "Finally explain the roadmap: Our next milestone is to integrate real specialist models, preserve original GeoTIFF uploads for live temporal analysis, and containerize the entire stack. The current architecture already isolates these concerns, reducing integration risk.")

    heading(doc, "16 Team operating checklist", 1)
    add_bullets(doc, [
        "Do not change the AnalysisResult contract without coordinating frontend and backend owners.",
        "Model outputs should be written to paths that the backend can expose as evidence artifacts.",
        "Preprocessing changes must preserve the task specific image count and output path guarantees.",
        "Agent changes must keep task names exactly vqa, grounding, change_detection, and optical_sar unless all consumers are updated together.",
        "Before merging, run the routing tests and add task focused tests for any changed integration boundary.",
        "For a live demonstration, use the same Python environment for Streamlit, FastAPI, and LangGraph dependencies.",
    ])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
